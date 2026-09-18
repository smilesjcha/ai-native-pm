"""Slide layouts — design/ppt-design-system.md 기준.

seoul-ai-foundation build_pptx 포팅 + KEDI 덱 벤치마크 레이아웃
(cover, navy_divider, toc, profile, closing) 추가.
"""

import os

from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

import theme as T
import shapes as S

ASSETS = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "assets"))

FOOTER_LEFT = "AI 시대의 업무 경쟁력과 생산성 향상 · 2026 청년 여성 밸런스 워크숍"
TOTAL = 53  # build.py에서 len(SLIDES)로 덮어씀


def _add_blank_slide(prs, bg=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg or T.WHITE
    return slide


def _add_title(slide, text, *, top=None, left=None, width=None, color=None,
               size=None):
    return S.add_textbox(
        slide,
        left or T.CONTENT_LEFT,
        top or T.CONTENT_TOP,
        width or T.CONTENT_WIDTH,
        Inches(0.7),
        text,
        size=size or T.PT_TITLE,
        bold=True,
        color=color or T.BLACK,
    )


def _title_rule(slide):
    S.add_thin_line(slide, T.CONTENT_LEFT, Inches(1.4), Inches(0.8),
                    color=T.DEEP_BLUE, weight=Pt(2.5))


def _add_footer(slide, slide_no, *, on_navy=False):
    color = T.BLUE_10 if on_navy else T.GRAY_80
    if not on_navy:
        S.add_textbox(slide, Inches(0.6), T.SLIDE_H - Inches(0.4), Inches(8),
                      Inches(0.3), FOOTER_LEFT,
                      size=T.PT_CAPTION, color=color)
    S.add_textbox(slide, T.SLIDE_W - Inches(1.6), T.SLIDE_H - Inches(0.4),
                  Inches(1.0), Inches(0.3),
                  f"{slide_no:02d} / {TOTAL:02d}",
                  size=T.PT_CAPTION, color=color, align=PP_ALIGN.RIGHT)


def _attach_note(slide, data):
    parts = []
    if data.get("note"):
        parts.append(data["note"].strip())
    if data.get("speak_seconds"):
        parts.append(f"[권장 발화 시간: {data['speak_seconds']}]")
    S.set_slide_notes(slide, "\n\n".join(parts))


# ---------------------------------------------------------------- Cover (KEDI S1 벤치마크)

def make_cover(prs, data):
    slide = _add_blank_slide(prs)
    S.add_textbox(slide, Inches(0.7), Inches(0.55), Inches(10), Inches(0.4),
                  "2026 청년 여성 밸런스 워크숍 (WORK-LIFE.zip)",
                  size=Pt(14), color=T.DEEP_BLUE, bold=True)
    S.add_thin_line(slide, Inches(0.7), Inches(1.05), Inches(1.6),
                    color=T.DEEP_BLUE, weight=Pt(2))
    # 네이비 패널 (풀 블리드)
    panel_top = Inches(2.05)
    panel_h = Inches(2.5)
    S.add_rect(slide, Emu(0), panel_top, T.SLIDE_W, panel_h, fill=T.BLUE_90)
    S.add_textbox(slide, Inches(0.7), panel_top + Inches(0.55),
                  T.SLIDE_W - Inches(1.4), Inches(1.0),
                  data.get("title", ""),
                  size=T.PT_COVER_TITLE, bold=True, color=T.WHITE,
                  align=PP_ALIGN.CENTER)
    S.add_textbox(slide, Inches(0.7), panel_top + Inches(1.62),
                  T.SLIDE_W - Inches(1.4), Inches(0.5),
                  data.get("subtitle", ""),
                  size=Pt(18), color=T.BLUE_10, align=PP_ALIGN.CENTER)
    S.add_textbox(slide, Inches(0.7), Inches(5.0), T.SLIDE_W - Inches(1.4),
                  Inches(0.45),
                  "2026. 07. 23. (목) 13:00 – 17:00 · 온라인(Zoom)",
                  size=Pt(16), color=T.GRAY_80, align=PP_ALIGN.CENTER)
    S.add_thin_line(slide, Inches(0.7), Inches(6.35), Inches(11.93),
                    color=T.GRAY_20, weight=Pt(0.5))
    S.add_textbox(slide, Inches(0.7), Inches(6.55), T.SLIDE_W - Inches(1.4),
                  Inches(0.4),
                  "차성재 · 무신사 Agentic AI PM · 서울시립대 / 아주대 AI 부문 겸임교수",
                  size=T.PT_BODY_SM, color=T.GRAY_80, align=PP_ALIGN.CENTER)
    _attach_note(slide, data)
    return slide


# ---------------------------------------------------------------- Navy divider (KEDI 시그니처)

def make_navy_divider(prs, data):
    slide = _add_blank_slide(prs, bg=T.BLUE_90)
    kicker = data.get("kicker", "")
    if kicker:
        S.add_textbox(slide, Inches(0.7), Inches(2.55), T.SLIDE_W - Inches(1.4),
                      Inches(0.45), kicker,
                      size=Pt(16), bold=True, color=T.BLUE_70_LIGHT,
                      align=PP_ALIGN.CENTER)
    S.add_textbox(slide, Inches(0.7), Inches(3.1), T.SLIDE_W - Inches(1.4),
                  Inches(1.1), data.get("title", ""),
                  size=T.PT_SECTION_TITLE, bold=True, color=T.WHITE,
                  align=PP_ALIGN.CENTER, line_spacing=1.2)
    sub = data.get("sub", "")
    if sub:
        S.add_textbox(slide, Inches(0.7), Inches(4.45), T.SLIDE_W - Inches(1.4),
                      Inches(0.6), sub,
                      size=Pt(16), color=T.BLUE_10, align=PP_ALIGN.CENTER,
                      line_spacing=1.3)
    _attach_note(slide, data)
    _add_footer(slide, data["no"], on_navy=True)
    return slide


# ---------------------------------------------------------------- TOC (KEDI S2 벤치마크)

def make_toc(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data.get("title", "목차"))
    _title_rule(slide)
    col_w = (T.CONTENT_WIDTH - Inches(0.5)) / 2
    groups = data.get("groups", [])
    for gi, group in enumerate(groups[:2]):
        left = T.CONTENT_LEFT + (col_w + Inches(0.5)) * gi
        S.add_filled_text_rect(
            slide, left, Inches(1.9), col_w, Inches(0.55),
            group["title"], fill=T.DEEP_BLUE, line=None, text_color=T.WHITE,
            size=T.PT_BODY, bold=True, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.MIDDLE)
        items = group.get("items", [])
        tb_lines = [f"{i+1}.  {it}" for i, it in enumerate(items)]
        S.add_multi_line_textbox(
            slide, left + Inches(0.1), Inches(2.75), col_w - Inches(0.2),
            Inches(4.2), tb_lines,
            size=T.PT_BODY, color=T.BLACK, line_spacing=1.75)
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Profile (KEDI S5-6 벤치마크)

def make_profile(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data.get("title", "강사 소개"))
    _title_rule(slide)
    photo = os.path.join(ASSETS, "profile_photo.jpg")
    if os.path.exists(photo):
        slide.shapes.add_picture(photo, T.CONTENT_LEFT, Inches(1.85),
                                 width=Inches(4.6))  # 759x437 → h≈2.65
    book = os.path.join(ASSETS, "book_cover.png")
    if os.path.exists(book):
        slide.shapes.add_picture(book, T.CONTENT_LEFT, Inches(4.8),
                                 height=Inches(2.05))  # 1016x1317 → w≈1.58
    S.add_multi_line_textbox(
        slide, T.CONTENT_LEFT + Inches(1.85), Inches(4.95), Inches(2.9),
        Inches(1.9),
        data.get("book_lines", []),
        size=T.PT_CAPTION, color=T.GRAY_80, line_spacing=1.4)
    right_left = T.CONTENT_LEFT + Inches(5.15)
    right_w = T.CONTENT_WIDTH - Inches(5.15)
    S.add_textbox(slide, right_left, Inches(1.85), right_w, Inches(0.5),
                  data.get("name_line", ""),
                  size=Pt(24), bold=True, color=T.BLACK)
    career_top = Inches(2.55)
    role_lines = data.get("role_lines", [])
    if role_lines:
        S.add_multi_line_textbox(
            slide, right_left, Inches(2.45), right_w, Inches(0.9),
            role_lines, size=Pt(17), bold=True, color=T.DEEP_BLUE,
            line_spacing=1.4)
        career_top = Inches(3.4)
    S.add_multi_line_textbox(
        slide, right_left, career_top, right_w, Inches(3.5),
        data.get("career_lines", []),
        size=T.PT_BODY_SM, color=T.BLACK, line_spacing=1.5)
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Closing (KEDI S95·100 벤치마크)

def make_closing(prs, data):
    slide = _add_blank_slide(prs, bg=T.BLUE_90)
    S.add_textbox(slide, Inches(0.7), Inches(2.3), T.SLIDE_W - Inches(1.4),
                  Inches(0.9), data.get("title", "감사합니다"),
                  size=T.PT_COVER_TITLE, bold=True, color=T.WHITE,
                  align=PP_ALIGN.CENTER)
    sub = data.get("sub", "")
    if sub:
        S.add_textbox(slide, Inches(0.7), Inches(3.45), T.SLIDE_W - Inches(1.4),
                      Inches(0.6), sub,
                      size=Pt(18), color=T.BLUE_10, align=PP_ALIGN.CENTER)
    lines = data.get("contact_lines", [])
    if lines:
        S.add_multi_line_textbox(
            slide, Inches(0.7), Inches(4.6), T.SLIDE_W - Inches(1.4),
            Inches(1.6), lines,
            size=T.PT_BODY_SM, color=T.BLUE_10, line_spacing=1.6,
            align=PP_ALIGN.CENTER)
    _attach_note(slide, data)
    _add_footer(slide, data["no"], on_navy=True)
    return slide


# ---------------------------------------------------------------- Concept

def make_concept(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    body = data.get("body", [])
    lead = data.get("lead")
    body_top = Inches(1.9)
    if lead:
        S.add_textbox(slide, T.CONTENT_LEFT, Inches(1.7),
                      T.CONTENT_WIDTH, Inches(1.0),
                      lead, size=T.PT_LEAD, color=T.DEEP_BLUE,
                      bold=True, line_spacing=1.3)
        body_top = Inches(2.9)
    if body:
        S.add_multi_line_textbox(
            slide, T.CONTENT_LEFT, body_top,
            T.CONTENT_WIDTH, Inches(4.5),
            body, size=T.PT_BODY, color=T.BLACK,
            line_spacing=1.55, bullet=True,
        )
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- One-liner

def make_one_liner(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    msg = data.get("body", [""])[0] if data.get("body") else ""
    S.add_rect(slide, T.CONTENT_LEFT, Inches(2.6), Emu(45720),
               Inches(2.5), fill=T.DEEP_BLUE)
    S.add_textbox(slide, T.CONTENT_LEFT + Inches(0.35), Inches(2.6),
                  T.CONTENT_WIDTH - Inches(0.35), Inches(2.5), msg,
                  size=Pt(27), color=T.BLACK, bold=True,
                  anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.4)
    sub = data.get("sub_lines", [])
    if sub:
        S.add_multi_line_textbox(slide, T.CONTENT_LEFT + Inches(0.35),
                                 Inches(5.5), T.CONTENT_WIDTH - Inches(0.35),
                                 Inches(1.2), sub,
                                 size=T.PT_BODY, color=T.GRAY_80,
                                 line_spacing=1.5)
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Process

def make_process(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    items = data.get("flow_items") or []
    if items:
        S.add_horizontal_flow(slide, T.CONTENT_LEFT, Inches(2.3),
                              T.CONTENT_WIDTH, Inches(1.4), items,
                              emphasize_index=data.get("emphasize_index"))
    body = data.get("body", [])
    if body:
        S.add_multi_line_textbox(
            slide, T.CONTENT_LEFT, Inches(4.3),
            T.CONTENT_WIDTH, Inches(2.6),
            body, size=T.PT_BODY, color=T.BLACK,
            line_spacing=1.5, bullet=True,
        )
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Document (좌 설명 + 우 박스)

def make_document(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    body = data.get("body", [])
    badge_text = data.get("badge")
    body_top = Inches(1.9)
    if badge_text:
        S.add_badge(slide, T.CONTENT_LEFT, Inches(1.85), Inches(1.6),
                    Inches(0.42), badge_text)
        body_top = Inches(2.5)
    S.add_multi_line_textbox(
        slide, T.CONTENT_LEFT, body_top,
        Inches(6.0), Inches(5.0),
        body, size=T.PT_BODY, color=T.BLACK,
        line_spacing=1.55, bullet=True,
    )
    box_left = T.CONTENT_LEFT + Inches(6.7)
    box_w = T.CONTENT_WIDTH - Inches(6.7)
    box_top = Inches(1.9)
    box_h = Inches(5.0)
    S.add_filled_text_rect(
        slide, box_left, box_top, box_w, Inches(0.55),
        data.get("box_title", "안내"),
        fill=T.DEEP_BLUE, line=None, text_color=T.WHITE,
        size=T.PT_BODY_SM, bold=True, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.MIDDLE)
    S.add_rect(slide, box_left, box_top + Inches(0.55), box_w,
               box_h - Inches(0.55), fill=T.BLUE_10, line=T.DEEP_BLUE,
               line_width=Pt(0.75))
    pad = Inches(0.25)
    S.add_multi_line_textbox(
        slide, box_left + pad, box_top + Inches(0.55) + pad,
        box_w - 2 * pad, box_h - Inches(0.55) - 2 * pad,
        data.get("box_lines", []), size=T.PT_BODY_SM, color=T.BLACK,
        line_spacing=1.4, bullet=False,
    )
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Code demo (프롬프트 예시)

def make_code_demo(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    body = data.get("body", [])
    S.add_multi_line_textbox(
        slide, T.CONTENT_LEFT, Inches(1.9),
        Inches(5.4), Inches(5.0),
        body, size=T.PT_BODY, color=T.BLACK,
        line_spacing=1.55, bullet=True,
    )
    code_left = T.CONTENT_LEFT + Inches(6.0)
    code_w = T.CONTENT_WIDTH - Inches(6.0)
    S.add_code_block(slide, code_left, Inches(1.9), code_w, Inches(5.0),
                     data.get("code_lines", []),
                     highlight_indices=data.get("code_highlight"))
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Review (checklist)

def make_review(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    lead = data.get("lead")
    items_top = Inches(1.9)
    if lead:
        S.add_callout_box(slide, T.CONTENT_LEFT, Inches(1.7),
                          T.CONTENT_WIDTH, Inches(0.8), lead,
                          size=T.PT_LEAD, bold=True)
        items_top = Inches(2.8)
    items = data.get("body", [])
    S.add_check_list(slide, T.CONTENT_LEFT + Inches(0.2), items_top,
                     T.CONTENT_WIDTH - Inches(0.4), Inches(4.0),
                     items, size=T.PT_BODY, line_spacing=1.7)
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Wrap-up

def make_wrap_up(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    body = data.get("body", [])
    S.add_multi_line_textbox(
        slide, T.CONTENT_LEFT, Inches(2.0),
        T.CONTENT_WIDTH, Inches(4.5),
        body, size=Pt(20), color=T.BLACK, bold=True,
        line_spacing=1.7, bullet=True,
    )
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Schedule table

def make_schedule(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    rows = data.get("schedule_rows", [])
    S.add_table(slide, T.CONTENT_LEFT, Inches(1.8),
                T.CONTENT_WIDTH, Inches(4.9),
                data.get("headers", ["시간", "세션", "핵심 내용", "산출물"]),
                rows,
                col_widths=data.get("col_widths",
                                    [Inches(1.9), Inches(2.4), Inches(5.7),
                                     Inches(2.1)]))
    note = data.get("body", [])
    if note:
        S.add_multi_line_textbox(slide, T.CONTENT_LEFT, Inches(6.8),
                                 T.CONTENT_WIDTH, Inches(0.4),
                                 note, size=T.PT_CAPTION, color=T.GRAY_80,
                                 line_spacing=1.3)
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Card grid

def make_card_grid(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    cards = data.get("cards", [])
    cols = data.get("card_cols", 2)
    rows = (len(cards) + cols - 1) // cols
    gap = Inches(0.25)
    body = data.get("body", [])
    bottom_reserve = Inches(0.6) if body else Inches(0)
    avail_h = Inches(4.85) - gap * (rows - 1) - bottom_reserve
    avail_w = T.CONTENT_WIDTH - gap * (cols - 1)
    card_w = avail_w / cols
    card_h = avail_h / rows
    top0 = Inches(1.85)
    for i, c in enumerate(cards):
        r = i // cols
        col = i % cols
        left = T.CONTENT_LEFT + (card_w + gap) * col
        top = top0 + (card_h + gap) * r
        S.add_card(slide, int(left), int(top), int(card_w), int(card_h),
                   c["title"], c.get("body", []),
                   accent=c.get("accent", False))
    if body:
        S.add_multi_line_textbox(slide, T.CONTENT_LEFT, Inches(6.72),
                                 T.CONTENT_WIDTH, Inches(0.4),
                                 body, size=T.PT_CAPTION, color=T.GRAY_80,
                                 line_spacing=1.3)
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Comparison

def make_comparison(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    columns = data.get("columns", [])
    n = len(columns)
    # 하단 캡션이 있으면 컬럼 박스를 짧게 잡아 겹침 방지
    body_h = Inches(3.9) if data.get("body") else Inches(4.3)
    if n:
        gap = Inches(0.3)
        col_w = (T.CONTENT_WIDTH - gap * (n - 1)) / n
        top = Inches(2.0)
        for i, col in enumerate(columns):
            left = T.CONTENT_LEFT + (col_w + gap) * i
            accent = col.get("accent", False)
            S.add_filled_text_rect(
                slide, int(left), top, int(col_w), Inches(0.6),
                col["title"],
                fill=T.DEEP_BLUE if accent else T.WHITE,
                line=None if accent else T.GRAY_20,
                text_color=T.WHITE if accent else T.BLACK,
                size=T.PT_BODY, bold=True, align=PP_ALIGN.CENTER)
            S.add_rect(slide, int(left), int(top + Inches(0.6)),
                       int(col_w), body_h,
                       fill=T.BLUE_10 if accent else T.GRAY_5,
                       line=T.DEEP_BLUE if accent else T.GRAY_20,
                       line_width=Pt(0.75))
            pad = Inches(0.22)
            S.add_multi_line_textbox(
                slide, int(left) + pad, int(top + Inches(0.6)) + pad,
                int(col_w) - 2 * pad, body_h - 2 * pad,
                col.get("body", []),
                size=T.PT_BODY_SM, color=T.BLACK,
                line_spacing=1.45, bullet=True)
    body = data.get("body", [])
    if body:
        S.add_multi_line_textbox(slide, T.CONTENT_LEFT, Inches(6.72),
                                 T.CONTENT_WIDTH, Inches(0.4),
                                 body, size=T.PT_CAPTION, color=T.GRAY_80,
                                 line_spacing=1.3)
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Shot (스크린샷 슬라이드)

def make_shot(prs, data):
    """스크린샷 중심 슬라이드.

    assets/screenshots/<image> 파일이 있으면 영역에 맞춰 자동 삽입,
    없으면 파일명이 적힌 플레이스홀더를 그린다 (파일 추가 후 재빌드하면 자동 반영).
    """
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    badge_text = data.get("badge")
    if badge_text:
        S.add_badge(slide, T.SLIDE_W - Inches(2.3), Inches(0.55),
                    Inches(1.7), Inches(0.42), badge_text)
    img_name = data.get("image", "")
    img_path = os.path.join(ASSETS, "screenshots", img_name)
    area_left = T.CONTENT_LEFT
    area_top = Inches(1.75)
    area_w = T.CONTENT_WIDTH
    area_h = Inches(4.55)
    caption = data.get("caption", "")
    if os.path.exists(img_path):
        try:
            from PIL import Image as _Img
            with _Img.open(img_path) as im:
                iw, ih = im.size
            scale = min(area_w / iw, area_h / ih)
            w = int(iw * scale)
            h = int(ih * scale)
            left = int(area_left + (area_w - w) / 2)
            top = int(area_top + (area_h - h) / 2)
            pic = slide.shapes.add_picture(img_path, left, top,
                                           width=w, height=h)
            pic.line.color.rgb = T.GRAY_20
            pic.line.width = Pt(0.75)
        except Exception:
            slide.shapes.add_picture(img_path, area_left, area_top,
                                     width=area_w)
    else:
        S.add_rect(slide, area_left, area_top, area_w, area_h,
                   fill=T.GRAY_5, line=T.GRAY_20, line_width=Pt(0.75))
        S.add_textbox(slide, area_left, area_top + Inches(1.7), area_w,
                      Inches(0.5),
                      f"스크린샷 자리 — assets/screenshots/{img_name}",
                      size=T.PT_BODY, color=T.GRAY_80, bold=True,
                      align=PP_ALIGN.CENTER)
        S.add_textbox(slide, area_left, area_top + Inches(2.3), area_w,
                      Inches(0.5),
                      "파일 저장 후 python3 build_pptx/build.py 재실행하면 자동 삽입됩니다",
                      size=T.PT_CAPTION, color=T.GRAY_80,
                      align=PP_ALIGN.CENTER)
    if caption:
        S.add_textbox(slide, area_left, Inches(6.45), area_w, Inches(0.5),
                      caption, size=T.PT_BODY_SM, color=T.GRAY_80,
                      align=PP_ALIGN.CENTER)
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide


# ---------------------------------------------------------------- Tree (folder)

def make_tree(prs, data):
    slide = _add_blank_slide(prs)
    _add_title(slide, data["title"])
    _title_rule(slide)
    body = data.get("body", [])
    S.add_multi_line_textbox(
        slide, T.CONTENT_LEFT, Inches(1.95),
        Inches(5.0), Inches(5.0),
        body, size=T.PT_BODY, color=T.BLACK,
        line_spacing=1.55, bullet=True,
    )
    tree_left = T.CONTENT_LEFT + Inches(5.6)
    tree_w = T.CONTENT_WIDTH - Inches(5.6)
    S.add_code_block(slide, tree_left, Inches(1.95), tree_w, Inches(5.0),
                     data.get("tree_lines", []),
                     highlight_indices=data.get("tree_highlight"))
    _attach_note(slide, data)
    _add_footer(slide, data["no"])
    return slide
